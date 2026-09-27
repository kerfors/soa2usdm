"""Apply REDACTIONS.json to the prompt and taxonomy and write the blind copies.

Fails if any logged span does not occur exactly once in its source, or if any study identifier
survives in the written instructions. Env: SOA2USDM_REPO, SOA2USDM_COLLECTIONS, SOA2USDM_BLIND.
"""
import json
import os
import re
import shutil
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
REPO = Path(os.environ["SOA2USDM_REPO"])
COLLECTIONS = Path(os.environ["SOA2USDM_COLLECTIONS"])
OUT = Path(os.environ["SOA2USDM_BLIND"]) / "_instructions"

# Names not derivable from the manifests: sponsors as spelled in prose, and the molecules.
EXTRA_TERMS = ["Lilly", "Novo", "Roche", "Tesaro", "Sanofi", "BMS", "AstraZeneca", "Amgen", "Alexion",
               "KalVista", "PrECOG", "niraparib", "isatuximab", "tirzepatide", "semaglutide",
               "bempegaldesleukin", "abemaciclib", "nivolumab", "mirikizumab", "CDISC_Pilot", "CDISC Pilot"]

log = json.loads((HERE / "REDACTIONS.json").read_text(encoding="utf-8"))

by_file = {}
for e in log["entries"]:
    by_file.setdefault(e["file"], []).append(e)

OUT.mkdir(parents=True, exist_ok=True)
for rel, entries in by_file.items():
    text = (REPO / rel).read_text(encoding="utf-8")
    for e in entries:
        n = text.count(e["find"])
        if n != 1:
            sys.exit(f"{e['id']}: span occurs {n}x in {rel} (expected 1) — update REDACTIONS.json")
        text = text.replace(e["find"], e["replace"])
    (OUT / Path(rel).name).write_text(text, encoding="utf-8")
    print(f"  {rel}: {len(entries)} span(s) redacted")
for rel in log["copied_unchanged"]:
    shutil.copyfile(REPO / rel, OUT / Path(rel).name)
    print(f"  {rel}: copied unchanged")

# Identifier scan over everything written.
terms = set(EXTRA_TERMS)
for book in COLLECTIONS.glob("*/studies_protocols.xlsx"):
    ws = openpyxl.load_workbook(book)["studies"]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        r = dict(zip(header, row))
        for k in ("nct_id", "study_code", "study_acronym"):
            if r.get(k):
                terms.add(str(r[k]))
        if r.get("d4k_folder"):
            terms.add(str(r["d4k_folder"]).split("_")[0])
terms.discard("Inv")   # d4k_folder prefix for 'Investigator'; not an identifier in prose
terms.discard("CDISC") # d4k_folder prefix of CDISC_Pilot; the standard's name, not a study (CDISC_Pilot/LZZT stay)
pats = [re.compile(r"NCT\d{8}")] + [re.compile(r"(?<![A-Za-z0-9])" + re.escape(t) + r"(?![A-Za-z0-9])") for t in terms]
hits = []
for f in sorted(OUT.iterdir()):
    for i, line in enumerate(f.read_text(encoding="utf-8").splitlines(), 1):
        for p in pats:
            for m in p.finditer(line):
                hits.append(f"{f.name}:{i}: {m.group(0)}")
print(f"  identifier scan: {len(terms)} terms + NCT pattern, {len(hits)} hit(s)")
for h in hits:
    print("   ", h)
sys.exit(1 if hits else 0)
